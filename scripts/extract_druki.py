import json, os, sys, subprocess
from multiprocessing import Pool

OUT = "/home/porwat/sejm-dump"
BASE = os.path.join(OUT, "pdfs")
PDFS = os.path.join(OUT, "druk_pdfs.txt")


def collect():
    files = []
    for root, _, names in os.walk(BASE):
        for n in names:
            if not n.endswith(".part"):
                files.append(os.path.join(root, n))
    return sorted(files)


def one(dest):
    ext = os.path.splitext(dest)[1].lower()
    try:
        if ext == ".pdf":
            r = subprocess.run(["pdftotext", "-layout", dest, "-"],
                               capture_output=True, timeout=180)
            if r.returncode != 0:
                return {"dest": dest, "text": "", "pages": None,
                        "err": "pdftotext rc=" + str(r.returncode)}
            return {"dest": dest, "text": r.stdout.decode("utf-8", "replace"),
                    "pages": None, "err": None}
        if ext in (".txt", ".xml", ".htm", ".html"):
            text = open(dest, "rb").read().decode("utf-8", "replace")
            return {"dest": dest, "text": text, "pages": None, "err": None}
        if ext in (".docx", ".docm"):
            from docx import Document
            d = Document(dest)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return {"dest": dest, "text": "\n".join(parts),
                    "pages": None, "err": None}
        if ext == ".rtf":
            from striprtf.striprtf import rtf_to_text
            raw_txt = open(dest, "r", encoding="latin-1", errors="ignore").read()
            return {"dest": dest, "text": rtf_to_text(raw_txt),
                    "pages": None, "err": None}
        if ext == ".odt":
            from odf.opendocument import load
            from odf.text import P
            doc = load(dest)
            paras = []
            for p in doc.text.getElementsByType(P):
                chunks = []
                for node in p.childNodes:
                    if hasattr(node, "data"):
                        chunks.append(str(node.data))
                paras.append("".join(chunks))
            return {"dest": dest, "text": "\n".join(paras),
                    "pages": None, "err": None}
        return {"dest": dest, "text": "", "pages": None, "err": "no-extractor:" + ext}
    except Exception as e:
        return {"dest": dest, "text": "", "pages": None, "err": str(e)[:150]}


if __name__ == "__main__":
    files = collect()
    print("pdf files:", len(files), flush=True)
    done = 0
    with open(os.path.join(OUT, "druk_texts.jsonl"), "w", encoding="utf-8") as out, \
            Pool(6) as pool:
        for res in pool.imap_unordered(one, files, chunksize=20):
            p = res["dest"]
            rel = os.path.relpath(p, BASE)  # term6/123/123.pdf
            parts = rel.split(os.sep)
            term = int(parts[0].replace("term", ""))
            number = parts[1]
            att = parts[2]
            row = {
                "term": term,
                "number": number,
                "attachment": att,
                "pages": res["pages"],
                "text_len": len(res["text"]),
                "error": res["err"],
                "text": res["text"],
            }
            out.write(json.dumps(row, ensure_ascii=False) + "\n")
            done += 1
            if done % 1000 == 0:
                print("extracted", done, flush=True)
    print("EXTRACT DONE", done, flush=True)

    from datasets import Dataset
    tok = open(os.path.expanduser("~/.hf-token")).read().strip()
    ds = Dataset.from_json(os.path.join(OUT, "druk_texts.jsonl"))
    ds.push_to_hub("patrykorwat/sejm-votings", config_name="druk_texts", token=tok)
    print("pushed druk_texts:", len(ds), flush=True)

    card_path = os.path.join(OUT, "dataset_card.md")
    s = open(card_path).read()
    if "config_name: druk_texts" not in s:
        j = s.index("\n---\n", 3)
        s = s[:j] + "\n- config_name: druk_texts\n  data_files: \"druk_texts/*.parquet\"" + s[j:]
        open(card_path, "w").write(s)
    from huggingface_hub import HfApi
    HfApi(token=tok).upload_file(
        path_or_fileobj=card_path, path_in_repo="README.md",
        repo_id="patrykorwat/sejm-votings", repo_type="dataset",
        commit_message="Add druk_texts config (PDF full text)")
    print("card updated", flush=True)
